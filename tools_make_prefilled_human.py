#!/usr/bin/env python3
from pathlib import Path

from openpyxl import Workbook, load_workbook


OLD_PATH = Path("/Users/sebbo/Desktop/Bigtech API script google.xlsx")
GRAPHS_PATH = Path("/Users/sebbo/Desktop/GB/BTC GRAPHS NEWSLETTER/BitcoinGraphs.xlsx")
OUT_PATH = Path("/Users/sebbo/Desktop/GB/Newsletter/newsletter_data_prefilled_human.xlsx")


def normalize(value):
    return "" if value is None else str(value).strip()


def cleaned_rows(ws):
    rows = []
    last_non_empty = -1
    for row in ws.iter_rows(values_only=True):
        values = list(row)
        while values and values[-1] in (None, ""):
            values.pop()
        if any(cell not in (None, "") for cell in values):
            last_non_empty = len(rows)
        rows.append(values)
    if last_non_empty < 0:
        return []
    rows = rows[: last_non_empty + 1]
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def headers_map(header_row):
    mapping = {}
    for index, value in enumerate(header_row):
        key = normalize(value).lower()
        if key:
            mapping[key] = index
    return mapping


def cell(row, mapping, key, default=""):
    index = mapping.get(key)
    if index is None or index >= len(row):
        return default
    value = row[index]
    return default if value is None else value


def copy_raw_sheet(dst_wb, name, src_ws):
    ws = dst_wb.create_sheet(name)
    for row in cleaned_rows(src_ws):
        ws.append(row)


def main():
    old_wb = load_workbook(OLD_PATH, data_only=True)
    graph_wb = load_workbook(GRAPHS_PATH, data_only=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "meta"
    for row in cleaned_rows(old_wb["meta"]):
        ws.append(row)

    copy_raw_sheet(wb, "points", old_wb["points"])

    gs = wb.create_sheet("graph_settings")
    gs.append(["graph_key", "show", "title", "comment", "comment_position", "top_n", "series_filter"])
    gs.append(["btc_price", "yes", "BTC Price", "", "below", "60", ""])
    gs.append(["liquidations", "yes", "Liquidations", "", "below", "6", ""])
    gs.append(["treasuries", "yes", "Treasuries (Top Holders)", "", "below", "6", ""])
    gs.append(["circulating_btc", "yes", "Circulating BTC", "", "below", "", ""])
    gs.append(["ownership", "yes", "Supply Ownership", "", "below", "8", ""])

    copy_raw_sheet(wb, "live_prices", old_wb["live_prices"])

    src = cleaned_rows(graph_wb["BTC Price"])
    ws = wb.create_sheet("BTC Price")
    ws.append(["date", "price", "asset", "show"])
    if src:
        mapping = headers_map(src[0])
        for row in src[1:]:
            price = cell(row, mapping, "price", "")
            if normalize(price) == "":
                continue
            ws.append([cell(row, mapping, "date", ""), price, normalize(cell(row, mapping, "asset", "BTC")) or "BTC", "yes"])

    src = cleaned_rows(graph_wb["Treasuries"])
    ws = wb.create_sheet("Treasuries")
    ws.append(["entity", "btc", "holder_group", "show"])
    if src:
        mapping = headers_map(src[0])
        for row in src[1:]:
            row_type = normalize(cell(row, mapping, "row_type", "")).lower()
            entity = normalize(cell(row, mapping, "entity", ""))
            btc = cell(row, mapping, "btc", "")
            if row_type and row_type != "entity":
                continue
            if not entity or normalize(btc) == "":
                continue
            ws.append([entity, btc, normalize(cell(row, mapping, "holder_group", "")), "yes"])

    src = cleaned_rows(graph_wb["Circulating BTC"])
    ws = wb.create_sheet("Circulating BTC")
    ws.append(["as_of_date", "circulating_supply_btc", "max_supply_btc", "note", "show"])
    if src:
        mapping = headers_map(src[0])
        for row in src[1:]:
            ws.append(
                [
                    cell(row, mapping, "as_of_date", ""),
                    cell(row, mapping, "circulating_supply_btc", ""),
                    cell(row, mapping, "max_supply_btc", ""),
                    cell(row, mapping, "note", ""),
                    "yes",
                ]
            )

    src = cleaned_rows(graph_wb["Liquidations"])
    ws = wb.create_sheet("Liquidations")
    ws.append(["label", "longs", "shorts", "total", "period_type", "show"])
    if src:
        mapping = headers_map(src[0])
        for row in src[1:]:
            label = normalize(cell(row, mapping, "label", ""))
            longs = cell(row, mapping, "longs", "")
            shorts = cell(row, mapping, "shorts", "")
            total = cell(row, mapping, "total", "")
            if normalize(total) == "":
                try:
                    total = float(longs or 0) + float(shorts or 0)
                except Exception:
                    total = ""
            if not label and normalize(total) == "":
                continue
            ws.append([label, longs, shorts, total, normalize(cell(row, mapping, "period_type", "")), "yes"])

    src = cleaned_rows(graph_wb["Distribution"])
    ws = wb.create_sheet("Distribution")
    ws.append(["category", "amount_btc", "percent", "color", "show"])
    if src:
        mapping = headers_map(src[0])
        for row in src[1:]:
            category = normalize(cell(row, mapping, "category", ""))
            amount = cell(row, mapping, "amount_btc", "")
            if not category and normalize(amount) == "":
                continue
            ws.append([category, amount, cell(row, mapping, "percent", ""), normalize(cell(row, mapping, "color", "")), "yes"])

    if "Macro Categories" in graph_wb.sheetnames:
        copy_raw_sheet(wb, "Macro Categories", graph_wb["Macro Categories"])

    wb.save(OUT_PATH)
    print(f"CREATED: {OUT_PATH}")
    print("SHEETS:", wb.sheetnames)


if __name__ == "__main__":
    main()
